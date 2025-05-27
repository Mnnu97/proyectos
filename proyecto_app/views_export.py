# proyecto_app/views_export.py

from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from .models import Proyecto
import openpyxl
from django.core.exceptions import PermissionDenied


@login_required
def exportar_proyectos(request):
    """
    Exporta los proyectos visibles al usuario actual a un archivo Excel (.xlsx)
    - Admin / Superuser → Todos los proyectos
    - Usuario común → Solo sus proyectos asignados
    """
    user = request.user

    # Filtramos según rol del usuario
    if user.groups.filter(name='Admin').exists() or user.is_superuser:
        proyectos = Proyecto.objects.all().order_by('nombre')
    else:
        proyectos = Proyecto.objects.filter(usuarios=user).order_by('nombre')

    # Preparar respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="listado_proyectos.xlsx"'

    # Crear libro de Excel
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Proyectos"

    # Encabezados
    worksheet.append(['Nombre', 'Descripción', 'Fecha Inicio', 'Fecha Fin', 'Estado'])

    # Datos de cada proyecto
    for p in proyectos:
        worksheet.append([
            p.nombre,
            p.descripcion or '',
            str(p.fecha_inicio) if p.fecha_inicio else '',
            str(p.fecha_fin) if p.fecha_fin else '',
            p.get_estado_display(),
        ])

    # Guardar y devolver el archivo
    workbook.save(response)
    return response


@login_required
def exportar_proyecto_individual(request, pk):
    """
    Exporta un único proyecto a un archivo Excel (.xlsx)
    """
    try:
        proyecto = Proyecto.objects.get(pk=pk)
    except Proyecto.DoesNotExist:
        raise PermissionDenied("Este proyecto no existe o no tienes acceso.")

    # Verificar permisos
    user = request.user
    if not (user.groups.filter(name='Admin').exists() or user.is_superuser or proyecto.usuarios.filter(pk=user.pk).exists()):
        raise PermissionDenied("No tienes permiso para ver este proyecto")

    # Preparar respuesta
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="proyecto_{proyecto.nombre}.xlsx"'

    # Crear archivo Excel
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = 'Proyecto'

    # Datos del proyecto
    worksheet.append(['Campo', 'Valor'])
    worksheet.append(['Nombre', proyecto.nombre])
    worksheet.append(['Descripción', proyecto.descripcion or ''])
    worksheet.append(['Fecha inicio', str(proyecto.fecha_inicio) if proyecto.fecha_inicio else ''])
    worksheet.append(['Fecha fin', str(proyecto.fecha_fin) if proyecto.fecha_fin else ''])
    worksheet.append(['Estado', proyecto.get_estado_display()])

    # Guardar y devolver respuesta
    workbook.save(response)
    return response


@login_required
def exportar_todos_los_proyectos(request):
    """
    Exporta TODOS los proyectos registrados (sin filtrar por usuario)
    Solo accesible para Admin o Superuser
    """
    user = request.user

    if not (user.groups.filter(name='Admin').exists() or user.is_superuser):
        raise PermissionDenied("Solo los administradores pueden exportar todos los proyectos")

    # Obtenemos todos los proyectos
    proyectos = Proyecto.objects.all().order_by('nombre')

    # Preparar respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="todos_los_proyectos.xlsx"'

    # Crear libro de Excel
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Proyectos"

    # Encabezados
    worksheet.append(['Nombre', 'Descripción', 'Fecha Inicio', 'Fecha Fin', 'Estado'])

    # Datos de cada proyecto
    for p in proyectos:
        worksheet.append([
            p.nombre,
            p.descripcion or '',
            str(p.fecha_inicio) if p.fecha_inicio else '',
            str(p.fecha_fin) if p.fecha_fin else '',
            p.get_estado_display(),
        ])

    # Guardar y devolver el archivo
    workbook.save(response)
    return response